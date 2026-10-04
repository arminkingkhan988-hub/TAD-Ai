import os
import io
import re
import ast
import csv
import base64
import math
import operator
import json
from urllib.parse import quote, urlparse

import requests
from bs4 import BeautifulSoup
from flask import Flask, request, jsonify, render_template_string, send_file
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


# =========================================================
# APP
# =========================================================

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 12 * 1024 * 1024

GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "").strip()

GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"

TEXT_MODEL = "openai/gpt-oss-20b"
VISION_MODEL = "qwen/qwen3.8-27b"


# =========================================================
# HELPERS
# =========================================================

def groq_request(messages, model=TEXT_MODEL, temperature=0.3):
    if not GROQ_API_KEY:
        return None, "GROQ_API_KEY is not configured."

    headers = {
        "Authorization": f"Bearer {GROQ_API_KEY}",
        "Content-Type": "application/json"
    }

    payload = {
        "model": model,
        "messages": messages,
        "temperature": temperature
    }

    try:
        response = requests.post(
            GROQ_URL,
            headers=headers,
            json=payload,
            timeout=55
        )
    except requests.RequestException as e:
        return None, f"Network error: {str(e)}"

    if response.status_code >= 400:
        try:
            details = response.json()
        except Exception:
            details = response.text

        return None, {
            "error": "Groq API error",
            "details": details,
            "status_code": response.status_code
        }

    try:
        data = response.json()
        answer = data["choices"][0]["message"]["content"]
        return answer, None
    except Exception as e:
        return None, f"Invalid Groq response: {str(e)}"


def clean_text(text, limit=30000):
    text = text or ""
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit]


def file_to_data_url(file_storage):
    raw = file_storage.read()

    if not raw:
        raise ValueError("Empty file.")

    mime = file_storage.mimetype or "image/jpeg"

    if not mime.startswith("image/"):
        mime = "image/jpeg"

    encoded = base64.b64encode(raw).decode("utf-8")

    return f"data:{mime};base64,{encoded}"


def safe_json_error(message, status=400):
    return jsonify({"error": message}), status


# =========================================================
# SAFE CALCULATOR
# =========================================================

BIN_OPS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.FloorDiv: operator.floordiv,
    ast.Mod: operator.mod,
    ast.Pow: operator.pow,
}

UNARY_OPS = {
    ast.UAdd: operator.pos,
    ast.USub: operator.neg,
}

SAFE_MATH = {
    "sqrt": math.sqrt,
    "sin": math.sin,
    "cos": math.cos,
    "tan": math.tan,
    "asin": math.asin,
    "acos": math.acos,
    "atan": math.atan,
    "log": math.log,
    "log10": math.log10,
    "exp": math.exp,
    "fabs": math.fabs,
    "floor": math.floor,
    "ceil": math.ceil,
    "factorial": math.factorial,
    "pi": math.pi,
    "e": math.e,
    "tau": math.tau,
}


def safe_eval(node):
    if isinstance(node, ast.Expression):
        return safe_eval(node.body)

    if isinstance(node, ast.Constant):
        if isinstance(node.value, (int, float)):
            return node.value
        raise ValueError("Invalid constant")

    if isinstance(node, ast.BinOp):
        op_type = type(node.op)

        if op_type not in BIN_OPS:
            raise ValueError("Operator not allowed")

        left = safe_eval(node.left)
        right = safe_eval(node.right)

        if op_type == ast.Pow and abs(right) > 100:
            raise ValueError("Power too large")

        return BIN_OPS[op_type](left, right)

    if isinstance(node, ast.UnaryOp):
        op_type = type(node.op)

        if op_type not in UNARY_OPS:
            raise ValueError("Unary operator not allowed")

        return UNARY_OPS[op_type](safe_eval(node.operand))

    if isinstance(node, ast.Name):
        if node.id in SAFE_MATH:
            return SAFE_MATH[node.id]

        raise ValueError("Unknown name")

    if isinstance(node, ast.Call):
        if not isinstance(node.func, ast.Name):
            raise ValueError("Function not allowed")

        name = node.func.id

        if name not in SAFE_MATH:
            raise ValueError("Function not allowed")

        fn = SAFE_MATH[name]

        args = [safe_eval(a) for a in node.args]

        return fn(*args)

    raise ValueError("Expression not allowed")


def calculate_expression(expression):
    expression = expression.strip()

    if len(expression) > 300:
        raise ValueError("Expression too long")

    expression = expression.replace("^", "**")

    tree = ast.parse(expression, mode="eval")

    return safe_eval(tree)


# =========================================================
# SYSTEM PROMPTS
# =========================================================

BASE_SYSTEM = """
You are MedAI, a helpful AI assistant.

You can answer in English, Pashto, Dari, or the user's language.

Medical safety:
- You provide educational information, not a definitive medical diagnosis.
- Never pretend to be a doctor.
- For emergencies or dangerous symptoms, advise urgent professional medical care.
- Do not encourage unsafe self-treatment.
- Ask for important missing information when appropriate.

Be accurate, clear, practical, and concise.
"""

MODE_PROMPTS = {
    "normal": """
Answer normally and helpfully.
""",

    "study": """
Study Mode:
Explain concepts step by step.
Use simple examples.
End with a short summary and optionally practice questions.
""",

    "coding": """
Coding Mode:
Give correct, practical code.
Explain important errors and implementation details.
Prefer complete working examples.
""",

    "writing": """
Writing Mode:
Help write, rewrite, improve, summarize, or structure text.
Preserve the user's intended meaning.
""",

    "quiz": """
Quiz Mode:
Create useful questions from the user's topic.
Do not reveal answers immediately unless requested.
""",

    "translate_en": """
Translation Mode:
Translate the user's content into natural English.
Return only the translation unless explanation is requested.
""",

    "translate_ps": """
Translation Mode:
Translate the user's content into natural Pashto.
Return only the translation unless explanation is requested.
""",

    "translate_fa": """
Translation Mode:
Translate the user's content into natural Dari.
Return only the translation unless explanation is requested.
"""
}


# =========================================================
# HTML / FRONTEND
# =========================================================

PAGE = r'''
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">

<title>MedAI</title>

<style>

* {
    box-sizing: border-box;
}

:root {
    --bg: #f5f7fb;
    --panel: #ffffff;
    --panel2: #eef2f7;
    --text: #111827;
    --muted: #667085;
    --border: #dfe4ea;
    --primary: #2563eb;
    --primary2: #1d4ed8;
    --danger: #dc2626;
    --shadow: 0 10px 30px rgba(0,0,0,.08);
}

body.dark {
    --bg: #0b1020;
    --panel: #111827;
    --panel2: #182033;
    --text: #f8fafc;
    --muted: #aab4c5;
    --border: #293548;
    --primary: #3b82f6;
    --primary2: #2563eb;
    --danger: #ef4444;
}

body {
    margin: 0;
    background: var(--bg);
    color: var(--text);
    font-family: Arial, Helvetica, sans-serif;
    min-height: 100vh;
}

button,
input,
textarea,
select {
    font: inherit;
}

button {
    cursor: pointer;
}

.app {
    display: flex;
    min-height: 100vh;
}

/* SIDEBAR */

.sidebar {
    width: 280px;
    background: var(--panel);
    border-right: 1px solid var(--border);
    padding: 16px;
    position: fixed;
    left: 0;
    top: 0;
    bottom: 0;
    overflow-y: auto;
    z-index: 50;
}

.logo {
    display: flex;
    align-items: center;
    gap: 10px;
    font-size: 22px;
    font-weight: 800;
    margin-bottom: 18px;
}

.logo-icon {
    width: 42px;
    height: 42px;
    border-radius: 13px;
    background: var(--primary);
    color: white;
    display: grid;
    place-items: center;
}

.side-btn {
    width: 100%;
    border: 1px solid var(--border);
    background: var(--panel);
    color: var(--text);
    padding: 11px 12px;
    border-radius: 10px;
    margin-bottom: 8px;
    text-align: left;
}

.side-btn:hover {
    background: var(--panel2);
}

.side-title {
    margin: 18px 5px 8px;
    color: var(--muted);
    font-size: 12px;
    font-weight: 700;
    text-transform: uppercase;
}

/* MAIN */

.main {
    margin-left: 280px;
    width: calc(100% - 280px);
    min-height: 100vh;
    display: flex;
    flex-direction: column;
}

.topbar {
    height: 64px;
    border-bottom: 1px solid var(--border);
    background: var(--panel);
    display: flex;
    align-items: center;
    gap: 10px;
    padding: 10px 18px;
    position: sticky;
    top: 0;
    z-index: 30;
}

.topbar h1 {
    font-size: 18px;
    margin: 0;
}

.icon-btn {
    width: 40px;
    height: 40px;
    border: 1px solid var(--border);
    background: var(--panel);
    color: var(--text);
    border-radius: 10px;
}

.mobile-menu {
    display: none;
}

/* CHAT */

.chat {
    flex: 1;
    padding: 20px;
    max-width: 1100px;
    width: 100%;
    margin: auto;
}

.welcome {
    text-align: center;
    padding: 60px 15px 30px;
}

.welcome h2 {
    font-size: 32px;
    margin: 0 0 10px;
}

.welcome p {
    color: var(--muted);
}

.messages {
    display: flex;
    flex-direction: column;
    gap: 16px;
    padding-bottom: 160px;
}

.message {
    display: flex;
    gap: 10px;
}

.message.user {
    justify-content: flex-end;
}

.avatar {
    width: 38px;
    height: 38px;
    min-width: 38px;
    border-radius: 50%;
    display: grid;
    place-items: center;
    background: var(--panel2);
}

.message.user .avatar {
    order: 2;
    background: var(--primary);
    color: white;
}

.bubble {
    max-width: min(850px, 88%);
    padding: 13px 15px;
    border: 1px solid var(--border);
    background: var(--panel);
    border-radius: 15px;
    line-height: 1.55;
    white-space: pre-wrap;
    overflow-wrap: anywhere;
    box-shadow: 0 3px 10px rgba(0,0,0,.03);
}

.message.user .bubble {
    background: var(--primary);
    color: white;
    border-color: var(--primary);
}

/* INPUT */

.composer-wrap {
    position: fixed;
    left: 280px;
    right: 0;
    bottom: 0;
    background: linear-gradient(transparent, var(--bg) 28%);
    padding: 25px 20px 18px;
    z-index: 25;
}

.composer {
    max-width: 1000px;
    margin: auto;
}

.tool-row {
    display: flex;
    gap: 7px;
    overflow-x: auto;
    margin-bottom: 8px;
    padding-bottom: 4px;
}

.tool {
    white-space: nowrap;
    border: 1px solid var(--border);
    background: var(--panel);
    color: var(--text);
    border-radius: 999px;
    padding: 7px 11px;
    font-size: 12px;
}

.tool:hover {
    background: var(--panel2);
}

.input-box {
    background: var(--panel);
    border: 1px solid var(--border);
    border-radius: 16px;
    padding: 10px;
    display: flex;
    gap: 8px;
    align-items: flex-end;
    box-shadow: var(--shadow);
}

textarea {
    flex: 1;
    min-height: 46px;
    max-height: 180px;
    resize: vertical;
    border: 0;
    outline: none;
    background: transparent;
    color: var(--text);
    padding: 10px;
}

.send {
    width: 46px;
    height: 46px;
    border: 0;
    border-radius: 12px;
    background: var(--primary);
    color: white;
    font-size: 18px;
}

.send:hover {
    background: var(--primary2);
}

.status {
    text-align: center;
    color: var(--muted);
    font-size: 12px;
    margin-top: 7px;
}

/* PANELS */

.panel {
    position: fixed;
    right: 20px;
    top: 75px;
    width: min(450px, calc(100vw - 30px));
    max-height: calc(100vh - 100px);
    overflow-y: auto;
    background: var(--panel);
    border: 1px solid var(--border);
    border-radius: 16px;
    box-shadow: var(--shadow);
    padding: 16px;
    z-index: 100;
    display: none;
}

.panel.show {
    display: block;
}

.panel h3 {
    margin-top: 0;
}

.close-panel {
    float: right;
    border: 0;
    background: transparent;
    color: var(--text);
    font-size: 22px;
}

.field {
    margin: 10px 0;
}

.field label {
    display: block;
    margin-bottom: 5px;
    color: var(--muted);
    font-size: 13px;
}

.field input,
.field textarea,
.field select {
    width: 100%;
    border: 1px solid var(--border);
    border-radius: 10px;
    background: var(--panel);
    color: var(--text);
    padding: 10px;
}

.action {
    border: 0;
    background: var(--primary);
    color: white;
    padding: 10px 14px;
    border-radius: 10px;
}

.secondary {
    border: 1px solid var(--border);
    background: var(--panel2);
    color: var(--text);
    padding: 10px 14px;
    border-radius: 10px;
}

.result {
    margin-top: 12px;
    padding: 10px;
    background: var(--panel2);
    border-radius: 10px;
    white-space: pre-wrap;
    overflow-wrap: anywhere;
}

.search-result {
    border: 1px solid var(--border);
    border-radius: 10px;
    padding: 10px;
    margin: 8px 0;
}

.search-result a {
    color: var(--primary);
    font-weight: 700;
}

.image-grid {
    display: grid;
    grid-template-columns: repeat(2, 1fr);
    gap: 8px;
    margin-top: 12px;
}

.image-grid img {
    width: 100%;
    height: 150px;
    object-fit: cover;
    border-radius: 9px;
    background: var(--panel2);
}

/* MODES */

.mode-select {
    margin-left: auto;
    border: 1px solid var(--border);
    border-radius: 9px;
    padding: 8px;
    background: var(--panel);
    color: var(--text);
}

/* LOADING */

.typing {
    display: inline-flex;
    gap: 4px;
}

.typing span {
    width: 7px;
    height: 7px;
    background: var(--muted);
    border-radius: 50%;
    animation: blink 1s infinite;
}

.typing span:nth-child(2) {
    animation-delay: .15s;
}

.typing span:nth-child(3) {
    animation-delay: .3s;
}

@keyframes blink {
    0%, 100% { opacity: .25; }
    50% { opacity: 1; }
}

/* MOBILE */

@media (max-width: 800px) {

    .sidebar {
        transform: translateX(-100%);
        transition: .25s;
        width: 270px;
    }

    .sidebar.open {
        transform: translateX(0);
    }

    .main {
        margin-left: 0;
        width: 100%;
    }

    .mobile-menu {
        display: block;
    }

    .composer-wrap {
        left: 0;
        padding: 20px 10px 10px;
    }

    .chat {
        padding: 12px;
    }

    .welcome {
        padding-top: 35px;
    }

    .welcome h2 {
        font-size: 25px;
    }

    .bubble {
        max-width: 90%;
    }

    .mode-select {
        max-width: 110px;
    }
}

</style>
</head>

<body>

<div class="app">

    <!-- SIDEBAR -->
    <aside class="sidebar" id="sidebar">

        <div class="logo">
            <div class="logo-icon">🩺</div>
            <span>MedAI</span>
        </div>

        <button class="side-btn" id="newChatBtn">🆕 New Chat</button>
        <button class="side-btn" id="historyBtn">💾 Chat History</button>
        <button class="side-btn" id="searchChatBtn">🔎 Search Chat</button>
        <button class="side-btn" id="exportBtn">📤 Export Chat</button>

        <div class="side-title">Tools</div>

        <button class="side-btn tool-open" data-panel="webPanel">🌐 Web Search</button>
        <button class="side-btn tool-open" data-panel="calcPanel">🧮 Calculator</button>
        <button class="side-btn tool-open" data-panel="imageSearchPanel">🖼️ Image Search</button>
        <button class="side-btn tool-open" data-panel="uploadPanel">👁️ Image Analysis / OCR</button>
        <button class="side-btn tool-open" data-panel="documentPanel">📄 PDF / DOCX / TXT</button>
        <button class="side-btn tool-open" data-panel="dataPanel">📊 CSV / Excel</button>
        <button class="side-btn tool-open" data-panel="urlPanel">🔗 URL Analysis</button>
        <button class="side-btn tool-open" data-panel="generatePanel">🎨 Image Generation</button>
        <button class="side-btn tool-open" data-panel="docxPanel">📥 Create DOCX</button>

        <div class="side-title">Modes</div>

        <button class="side-btn mode-btn" data-mode="study">🎓 Study Mode</button>
        <button class="side-btn mode-btn" data-mode="coding">💻 Coding Mode</button>
        <button class="side-btn mode-btn" data-mode="writing">✍️ Writing Mode</button>
        <button class="side-btn mode-btn" data-mode="quiz">❓ Quiz Mode</button>

        <div class="side-title">Appearance</div>

        <button class="side-btn" id="themeBtn">🌙 Dark / Light</button>

    </aside>


    <!-- MAIN -->
    <main class="main">

        <header class="topbar">

            <button class="icon-btn mobile-menu" id="mobileMenu">☰</button>

            <h1>MedAI Assistant</h1>

            <select class="mode-select" id="modeSelect">
                <option value="normal">💬 Normal</option>
                <option value="study">🎓 Study</option>
                <option value="coding">💻 Coding</option>
                <option value="writing">✍️ Writing</option>
                <option value="quiz">❓ Quiz</option>
                <option value="translate_en">🌍 English</option>
                <option value="translate_ps">🌍 Pashto</option>
                <option value="translate_fa">🌍 Dari</option>
            </select>

            <button class="icon-btn" id="voiceOutputBtn" title="Voice output">🔊</button>

        </header>


        <section class="chat">

            <div class="welcome" id="welcome">
                <h2>How can I help you?</h2>
                <p>
                    AI Chat • Search • Images • OCR • Documents • Coding • Study • Medical Safety
                </p>
            </div>

            <div class="messages" id="messages"></div>

        </section>

    </main>

</div>


<!-- COMPOSER -->

<div class="composer-wrap">

    <div class="composer">

        <div class="tool-row">

            <button class="tool tool-open" data-panel="webPanel">🌐 Search</button>
            <button class="tool tool-open" data-panel="calcPanel">🧮 Calc</button>
            <button class="tool tool-open" data-panel="imageSearchPanel">🖼️ Images</button>
            <button class="tool tool-open" data-panel="uploadPanel">👁️ Image</button>
            <button class="tool tool-open" data-panel="documentPanel">📄 Files</button>
            <button class="tool tool-open" data-panel="dataPanel">📊 Data</button>
            <button class="tool tool-open" data-panel="urlPanel">🔗 URL</button>
            <button class="tool tool-open" data-panel="generatePanel">🎨 Generate</button>

        </div>

        <div class="input-box">

            <button class="icon-btn" id="voiceInputBtn" title="Voice input">🎤</button>

            <textarea
                id="messageInput"
                placeholder="Message MedAI..."
            ></textarea>

            <button class="send" id="sendBtn">➤</button>

        </div>

        <div class="status" id="status">
            MedAI • Educational medical assistant
        </div>

    </div>

</div>


<!-- WEB SEARCH -->

<div class="panel" id="webPanel">

    <button class="close-panel">×</button>

    <h3>🌐 Web Search</h3>

    <div class="field">
        <label>Search query</label>
        <input id="webQuery" placeholder="Search the web...">
    </div>

    <button class="action" id="webSearchBtn">Search</button>

    <div id="webResults"></div>

</div>


<!-- CALCULATOR -->

<div class="panel" id="calcPanel">

    <button class="close-panel">×</button>

    <h3>🧮 Advanced Calculator</h3>

    <div class="field">
        <label>Expression</label>
        <input id="calcInput" placeholder="sqrt(25) + 10 * 2">
    </div>

    <button class="action" id="calcBtn">Calculate</button>

    <div class="result" id="calcResult"></div>

</div>


<!-- IMAGE SEARCH -->

<div class="panel" id="imageSearchPanel">

    <button class="close-panel">×</button>

    <h3>🖼️ Image Search</h3>

    <div class="field">
        <label>Image search</label>
        <input id="imageQuery" placeholder="Human heart, lungs, computer...">
    </div>

    <button class="action" id="imageSearchBtn">Search Images</button>

    <div class="image-grid" id="imageResults"></div>

</div>


<!-- IMAGE / OCR -->

<div class="panel" id="uploadPanel">

    <button class="close-panel">×</button>

    <h3>👁️ Image Analysis + OCR</h3>

    <div class="field">
        <label>Select image</label>
        <input type="file" id="imageFile" accept="image/*">
    </div>

    <button class="action" id="analyzeImageBtn">Analyze Image</button>
    <button class="secondary" id="ocrBtn">OCR</button>

    <div class="result" id="imageResult"></div>

</div>


<!-- DOCUMENTS -->

<div class="panel" id="documentPanel">

    <button class="close-panel">×</button>

    <h3>📄 Document Analysis</h3>

    <div class="field">
        <label>PDF / DOCX / TXT</label>
        <input type="file" id="documentFile" accept=".pdf,.docx,.txt">
    </div>

    <div class="field">
        <label>Instruction</label>
        <textarea id="documentInstruction" placeholder="Summarize this document..."></textarea>
    </div>

    <button class="action" id="documentBtn">Analyze Document</button>

    <div class="result" id="documentResult"></div>

</div>


<!-- DATA -->

<div class="panel" id="dataPanel">

    <button class="close-panel">×</button>

    <h3>📊 CSV / Excel Analysis</h3>

    <div class="field">
        <label>CSV / XLSX</label>
        <input type="file" id="dataFile" accept=".csv,.xlsx,.xls">
    </div>

    <div class="field">
        <label>Question</label>
        <textarea id="dataQuestion" placeholder="Analyze this data..."></textarea>
    </div>

    <button class="action" id="dataBtn">Analyze Data</button>

    <div class="result" id="dataResult"></div>

</div>


<!-- URL -->

<div class="panel" id="urlPanel">

    <button class="close-panel">×</button>

    <h3>🔗 URL / Website Analysis</h3>

    <div class="field">
        <label>Website URL</label>
        <input id="urlInput" placeholder="https://example.com">
    </div>

    <div class="field">
        <label>Question</label>
        <textarea id="urlQuestion" placeholder="Summarize this website..."></textarea>
    </div>

    <button class="action" id="urlBtn">Analyze Website</button>

    <div class="result" id="urlResult"></div>

</div>


<!-- IMAGE GENERATION -->

<div class="panel" id="generatePanel">

    <button class="close-panel">×</button>

    <h3>🎨 Image Generation</h3>

    <p style="color:var(--muted);font-size:13px;">
        Uses a free public image generation service.
    </p>

    <div class="field">
        <label>Prompt</label>
        <textarea id="generatePrompt" placeholder="A futuristic Afghan mountain city at sunset..."></textarea>
    </div>

    <button class="action" id="generateBtn">Generate</button>

    <div id="generatedImage"></div>

</div>


<!-- DOCX -->

<div class="panel" id="docxPanel">

    <button class="close-panel">×</button>

    <h3>📥 Create DOCX</h3>

    <div class="field">
        <label>Title</label>
        <input id="docxTitle" placeholder="My document">
    </div>

    <div class="field">
        <label>Content</label>
        <textarea id="docxContent" style="min-height:180px;" placeholder="Write document content..."></textarea>
    </div>

    <button class="action" id="createDocxBtn">Create DOCX</button>

    <div class="result" id="docxResult"></div>

</div>


<script>

const $ = (id) => document.getElementById(id);

let messages = [];
let currentMode = localStorage.getItem("medai_mode") || "normal";
let voiceEnabled = false;


/* ========================================================
   THEME
======================================================== */

if (localStorage.getItem("medai_theme") === "dark") {
    document.body.classList.add("dark");
}

$("themeBtn").addEventListener("click", () => {

    document.body.classList.toggle("dark");

    localStorage.setItem(
        "medai_theme",
        document.body.classList.contains("dark") ? "dark" : "light"
    );

});


/* ========================================================
   MOBILE
======================================================== */

$("mobileMenu").addEventListener("click", () => {
    $("sidebar").classList.toggle("open");
});


/* ========================================================
   PANELS
======================================================== */

function closePanels() {
    document.querySelectorAll(".panel").forEach(p => {
        p.classList.remove("show");
    });
}

document.querySelectorAll(".tool-open").forEach(btn => {

    btn.addEventListener("click", () => {

        closePanels();

        const panel = $(btn.dataset.panel);

        if (panel) {
            panel.classList.add("show");
        }

    });

});

document.querySelectorAll(".close-panel").forEach(btn => {

    btn.addEventListener("click", () => {
        btn.parentElement.classList.remove("show");
    });

});


/* ========================================================
   MODE
======================================================== */

$("modeSelect").value = currentMode;

function setMode(mode) {

    currentMode = mode;

    $("modeSelect").value = mode;

    localStorage.setItem("medai_mode", mode);

    $("status").textContent = "Mode: " + mode;

}

$("modeSelect").addEventListener("change", e => {
    setMode(e.target.value);
});

document.querySelectorAll(".mode-btn").forEach(btn => {

    btn.addEventListener("click", () => {
        setMode(btn.dataset.mode);
        $("sidebar").classList.remove("open");
    });

});


/* ========================================================
   CHAT STORAGE
======================================================== */

function saveHistory() {
    localStorage.setItem("medai_messages", JSON.stringify(messages));
}

function loadHistory() {

    try {
        const saved = JSON.parse(
            localStorage.getItem("medai_messages") || "[]"
        );

        if (Array.isArray(saved)) {
            messages = saved;
        }

    } catch {
        messages = [];
    }

}

function escapeText(text) {
    return String(text || "");
}


/* ========================================================
   RENDER CHAT
======================================================== */

function renderMessages() {

    const container = $("messages");

    container.innerHTML = "";

    if (messages.length === 0) {
        $("welcome").style.display = "block";
        return;
    }

    $("welcome").style.display = "none";

    messages.forEach((msg, index) => {

        const row = document.createElement("div");

        row.className =
            "message " +
            (msg.role === "user" ? "user" : "assistant");

        const avatar = document.createElement("div");

        avatar.className = "avatar";

        avatar.textContent =
            msg.role === "user" ? "👤" : "🤖";

        const bubble = document.createElement("div");

        bubble.className = "bubble";

        bubble.textContent = escapeText(msg.content);

        row.appendChild(avatar);
        row.appendChild(bubble);

        container.appendChild(row);

    });

    window.scrollTo({
        top: document.body.scrollHeight,
        behavior: "smooth"
    });

}


loadHistory();
renderMessages();


/* ========================================================
   NEW CHAT
======================================================== */

$("newChatBtn").addEventListener("click", () => {

    messages = [];

    localStorage.removeItem("medai_messages");

    renderMessages();

    $("messageInput").value = "";

    $("status").textContent = "New chat started.";

    $("sidebar").classList.remove("open");

});


/* ========================================================
   SEND CHAT
======================================================== */

async function sendMessage() {

    const input = $("messageInput");

    const text = input.value.trim();

    if (!text) {
        return;
    }

    input.value = "";

    messages.push({
        role: "user",
        content: text
    });

    saveHistory();
    renderMessages();

    $("status").textContent = "MedAI is thinking...";

    const loading = document.createElement("div");

    loading.className = "message assistant";

    loading.id = "loadingMessage";

    loading.innerHTML = `
        <div class="avatar">🤖</div>
        <div class="bubble">
            <div class="typing">
                <span></span><span></span><span></span>
            </div>
        </div>
    `;

    $("messages").appendChild(loading);

    try {

        const response = await fetch("/chat", {

            method: "POST",

            headers: {
                "Content-Type": "application/json"
            },

            body: JSON.stringify({
                message: text,
                history: messages.slice(-20),
                mode: currentMode
            })

        });

        const data = await response.json();

        const loadingElement = $("loadingMessage");

        if (loadingElement) {
            loadingElement.remove();
        }

        if (!response.ok || data.error) {
            throw new Error(
                data.details
                    ? JSON.stringify(data.details)
                    : data.error || "Request failed"
            );
        }

        messages.push({
            role: "assistant",
            content: data.answer || "No answer received."
        });

        saveHistory();
        renderMessages();

        $("status").textContent = "Ready";

        if (voiceEnabled && data.answer) {
            speak(data.answer);
        }

    } catch (error) {

        const loadingElement = $("loadingMessage");

        if (loadingElement) {
            loadingElement.remove();
        }

        messages.push({
            role: "assistant",
            content: "Error: " + error.message
        });

        saveHistory();
        renderMessages();

        $("status").textContent = "Error";

    }

}

$("sendBtn").addEventListener("click", sendMessage);

$("messageInput").addEventListener("keydown", e => {

    if (e.key === "Enter" && !e.shiftKey) {
        e.preventDefault();
        sendMessage();
    }

});


/* ========================================================
   VOICE OUTPUT
======================================================== */

function speak(text) {

    if (!("speechSynthesis" in window)) {
        alert("Voice output is not supported in this browser.");
        return;
    }

    window.speechSynthesis.cancel();

    const utterance =
        new SpeechSynthesisUtterance(text);

    utterance.rate = 1;
    utterance.pitch = 1;

    window.speechSynthesis.speak(utterance);
}

$("voiceOutputBtn").addEventListener("click", () => {

    voiceEnabled = !voiceEnabled;

    $("voiceOutputBtn").textContent =
        voiceEnabled ? "🔊" : "🔇";

});


/* ========================================================
   VOICE INPUT
======================================================== */

$("voiceInputBtn").addEventListener("click", () => {

    const SpeechRecognition =
        window.SpeechRecognition ||
        window.webkitSpeechRecognition;

    if (!SpeechRecognition) {

        alert(
            "Voice input is not supported in this browser. " +
            "Try Chrome or Edge."
        );

        return;
    }

    const recognition = new SpeechRecognition();

    recognition.lang = "en-US";
    recognition.interimResults = false;

    recognition.onstart = () => {
        $("status").textContent = "Listening...";
    };

    recognition.onresult = event => {

        const transcript =
            event.results[0][0].transcript;

        $("messageInput").value +=
            ($("messageInput").value ? " " : "") +
            transcript;

    };

    recognition.onerror = event => {
        $("status").textContent =
            "Voice error: " + event.error;
    };

    recognition.onend = () => {
        $("status").textContent = "Ready";
    };

    recognition.start();

});


/* ========================================================
   WEB SEARCH
======================================================== */

$("webSearchBtn").addEventListener("click", async () => {

    const query = $("webQuery").value.trim();

    if (!query) return;

    $("webResults").innerHTML = "Searching...";

    try {

        const response = await fetch(
            "/search?q=" + encodeURIComponent(query)
        );

        const data = await response.json();

        if (data.error) {
            throw new Error(data.error);
        }

        const results = data.results || [];

        if (!results.length) {
            $("webResults").textContent =
                "No results found.";
            return;
        }

        $("webResults").innerHTML = "";

        results.forEach(item => {

            const div = document.createElement("div");

            div.className = "search-result";

            const a = document.createElement("a");

            a.href = item.url;
            a.target = "_blank";
            a.rel = "noopener noreferrer";
            a.textContent = item.title || item.url;

            const p = document.createElement("p");

            p.textContent = item.snippet || "";

            div.appendChild(a);
            div.appendChild(p);

            $("webResults").appendChild(div);

        });

    } catch (error) {

        $("webResults").textContent =
            "Error: " + error.message;

    }

});


/* ========================================================
   CALCULATOR
======================================================== */

$("calcBtn").addEventListener("click", async () => {

    const expression = $("calcInput").value.trim();

    if (!expression) return;

    $("calcResult").textContent = "Calculating...";

    try {

        const response = await fetch("/calculate", {

            method: "POST",

            headers: {
                "Content-Type": "application/json"
            },

            body: JSON.stringify({
                expression
            })

        });

        const data = await response.json();

        if (data.error) {
            throw new Error(data.error);
        }

        $("calcResult").textContent =
            String(data.result);

    } catch (error) {

        $("calcResult").textContent =
            "Error: " + error.message;

    }

});


/* ========================================================
   IMAGE SEARCH
======================================================== */

$("imageSearchBtn").addEventListener("click", async () => {

    const query = $("imageQuery").value.trim();

    if (!query) return;

    $("imageResults").innerHTML = "Searching...";

    try {

        const response = await fetch(
            "/images?q=" + encodeURIComponent(query)
        );

        const data = await response.json();

        if (data.error) {
            throw new Error(data.error);
        }

        $("imageResults").innerHTML = "";

        (data.images || []).forEach(item => {

            const img = document.createElement("img");

            img.src = item.url;
            img.alt = item.title || query;
            img.loading = "lazy";

            img.addEventListener("click", () => {
                window.open(item.page || item.url, "_blank");
            });

            $("imageResults").appendChild(img);

        });

    } catch (error) {

        $("imageResults").textContent =
            "Error: " + error.message;

    }

});


/* ========================================================
   IMAGE ANALYSIS
======================================================== */

async function sendImage(mode) {

    const file = $("imageFile").files[0];

    if (!file) {
        alert("Please select an image.");
        return;
    }

    const form = new FormData();

    form.append("image", file);
    form.append("mode", mode);

    $("imageResult").textContent =
        mode === "ocr"
            ? "Reading image..."
            : "Analyzing image...";

    try {

        const response = await fetch("/analyze-image", {
            method: "POST",
            body: form
        });

        const data = await response.json();

        if (!response.ok || data.error) {
            throw new Error(
                data.details
                    ? JSON.stringify(data.details)
                    : data.error
            );
        }

        $("imageResult").textContent =
            data.answer || "No result.";

    } catch (error) {

        $("imageResult").textContent =
            "Error: " + error.message;

    }

}

$("analyzeImageBtn").addEventListener(
    "click",
    () => sendImage("analysis")
);

$("ocrBtn").addEventListener(
    "click",
    () => sendImage("ocr")
);


/* ========================================================
   DOCUMENT
======================================================== */

$("documentBtn").addEventListener("click", async () => {

    const file = $("documentFile").files[0];

    if (!file) {
        alert("Select a document.");
        return;
    }

    const instruction =
        $("documentInstruction").value.trim() ||
        "Summarize and explain this document.";

    const form = new FormData();

    form.append("file", file);
    form.append("instruction", instruction);

    $("documentResult").textContent =
        "Analyzing document...";

    try {

        const response = await fetch(
            "/analyze-document",
            {
                method: "POST",
                body: form
            }
        );

        const data = await response.json();

        if (!response.ok || data.error) {
            throw new Error(
                data.details
                    ? JSON.stringify(data.details)
                    : data.error
            );
        }

        $("documentResult").textContent =
            data.answer || "No result.";

    } catch (error) {

        $("documentResult").textContent =
            "Error: " + error.message;

    }

});


/* ========================================================
   DATA
======================================================== */

$("dataBtn").addEventListener("click", async () => {

    const file = $("dataFile").files[0];

    if (!file) {
        alert("Select CSV or Excel file.");
        return;
    }

    const question =
        $("dataQuestion").value.trim() ||
        "Analyze this dataset and provide useful insights.";

    const form = new FormData();

    form.append("file", file);
    form.append("question", question);

    $("dataResult").textContent =
        "Analyzing data...";

    try {

        const response = await fetch(
            "/analyze-data",
            {
                method: "POST",
                body: form
            }
        );

        const data = await response.json();

        if (!response.ok || data.error) {
            throw new Error(
                data.details
                    ? JSON.stringify(data.details)
                    : data.error
            );
        }

        $("dataResult").textContent =
            data.answer || "No result.";

    } catch (error) {

        $("dataResult").textContent =
            "Error: " + error.message;

    }

});


/* ========================================================
   URL ANALYSIS
======================================================== */

$("urlBtn").addEventListener("click", async () => {

    const url = $("urlInput").value.trim();

    if (!url) return;

    const question =
        $("urlQuestion").value.trim() ||
        "Summarize this website.";

    $("urlResult").textContent =
        "Opening website...";

    try {

        const response = await fetch("/analyze-url", {

            method: "POST",

            headers: {
                "Content-Type": "application/json"
            },

            body: JSON.stringify({
                url,
                question
            })

        });

        const data = await response.json();

        if (!response.ok || data.error) {
            throw new Error(
                data.details
                    ? JSON.stringify(data.details)
                    : data.error
            );
        }

        $("urlResult").textContent =
            data.answer || "No result.";

    } catch (error) {

        $("urlResult").textContent =
            "Error: " + error.message;

    }

});


/* ========================================================
   IMAGE GENERATION
======================================================== */

$("generateBtn").addEventListener("click", async () => {

    const prompt =
        $("generatePrompt").value.trim();

    if (!prompt) return;

    $("generatedImage").innerHTML =
        "Generating...";

    try {

        const response = await fetch(
            "/generate-image",
            {
                method: "POST",
                headers: {
                    "Content-Type": "application/json"
                },
                body: JSON.stringify({
                    prompt
                })
            }
        );

        const data = await response.json();

        if (!response.ok || data.error) {
            throw new Error(data.error);
        }

        $("generatedImage").innerHTML = "";

        const img = document.createElement("img");

        img.src = data.url;

        img.style.width = "100%";
        img.style.borderRadius = "12px";
        img.style.marginTop = "12px";

        $("generatedImage").appendChild(img);

        const link = document.createElement("a");

        link.href = data.url;
        link.target = "_blank";
        link.textContent = "Open generated image";

        link.style.display = "block";
        link.style.marginTop = "8px";
        link.style.color = "var(--primary)";

        $("generatedImage").appendChild(link);

    } catch (error) {

        $("generatedImage").textContent =
            "Error: " + error.message;

    }

});


/* ========================================================
   DOCX
======================================================== */

$("createDocxBtn").addEventListener("click", async () => {

    const title =
        $("docxTitle").value.trim() ||
        "MedAI Document";

    const content =
        $("docxContent").value.trim();

    if (!content) {
        alert("Write some content first.");
        return;
    }

    $("docxResult").textContent =
        "Creating DOCX...";

    try {

        const response = await fetch(
            "/create-docx",
            {
                method: "POST",
                headers: {
                    "Content-Type": "application/json"
                },
                body: JSON.stringify({
                    title,
                    content
                })
            }
        );

        if (!response.ok) {

            const data = await response.json();

            throw new Error(
                data.error || "DOCX creation failed."
            );
        }

        const blob = await response.blob();

        const url =
            window.URL.createObjectURL(blob);

        const a =
            document.createElement("a");

        a.href = url;
        a.download =
            title.replace(/[^a-z0-9_-]/gi, "_") +
            ".docx";

        document.body.appendChild(a);

        a.click();

        a.remove();

        window.URL.revokeObjectURL(url);

        $("docxResult").textContent =
            "DOCX created successfully.";

    } catch (error) {

        $("docxResult").textContent =
            "Error: " + error.message;

    }

});


/* ========================================================
   CHAT SEARCH
======================================================== */

$("searchChatBtn").addEventListener("click", () => {

    const query = prompt("Search your local chat history:");

    if (!query) return;

    const q = query.toLowerCase();

    const matches = messages.filter(m =>
        String(m.content || "")
            .toLowerCase()
            .includes(q)
    );

    if (!matches.length) {

        alert("No matching messages found.");

        return;
    }

    let text = "";

    matches.forEach(m => {

        text +=
            (m.role === "user" ? "You: " : "MedAI: ") +
            m.content +
            "\n\n";

    });

    alert(text);

});


/* ========================================================
   HISTORY
======================================================== */

$("historyBtn").addEventListener("click", () => {

    if (!messages.length) {
        alert("No local chat history.");
        return;
    }

    alert(
        "Local history contains " +
        messages.length +
        " messages."
    );

});


/* ========================================================
   EXPORT CHAT
======================================================== */

$("exportBtn").addEventListener("click", () => {

    if (!messages.length) {
        alert("No chat to export.");
        return;
    }

    let text =
        "MedAI Chat Export\n" +
        "=================\n\n";

    messages.forEach(m => {

        text +=
            (m.role === "user" ? "USER" : "MEDAI") +
            ":\n" +
            m.content +
            "\n\n";

    });

    const blob = new Blob(
        [text],
        { type: "text/plain;charset=utf-8" }
    );

    const url =
        URL.createObjectURL(blob);

    const a =
        document.createElement("a");

    a.href = url;
    a.download = "medai-chat.txt";

    document.body.appendChild(a);

    a.click();

    a.remove();

    URL.revokeObjectURL(url);

});


/* ========================================================
   INITIAL
======================================================== */

$("status").textContent =
    "MedAI ready • Mode: " + currentMode;

</script>

</body>
</html>
'''


# =========================================================
# ROUTES
# =========================================================

@app.route("/", methods=["GET"])
def home():
    return render_template_string(PAGE)


@app.route("/health", methods=["GET"])
def health():
    return jsonify({
        "service": "MedAI",
        "status": "ok"
    })


# =========================================================
# CHAT
# =========================================================

@app.route("/chat", methods=["POST"])
def chat():

    data = request.get_json(silent=True) or {}

    user_message = str(
        data.get("message", "")
    ).strip()

    history = data.get("history", [])

    mode = str(
        data.get("mode", "normal")
    )

    if not user_message:
        return safe_json_error(
            "Message is required."
        )

    if mode not in MODE_PROMPTS:
        mode = "normal"

    messages_for_groq = [
        {
            "role": "system",
            "content":
                BASE_SYSTEM +
                "\n\n" +
                MODE_PROMPTS[mode]
        }
    ]

    if isinstance(history, list):

        for item in history[-20:]:

            if not isinstance(item, dict):
                continue

            role = item.get("role")

            content = item.get("content")

            if role not in ["user", "assistant"]:
                continue

            if not isinstance(content, str):
                continue

            messages_for_groq.append({
                "role": role,
                "content": content[:12000]
            })

    answer, error = groq_request(
        messages_for_groq,
        TEXT_MODEL,
        0.3
    )

    if error:
        if isinstance(error, dict):
            return jsonify(error), error.get(
                "status_code", 500
            )

        return jsonify({
            "error": error
        }), 500

    return jsonify({
        "answer": answer
    })


# =========================================================
# IMAGE SEARCH - WIKIMEDIA
# =========================================================

@app.route("/images", methods=["GET"])
def images():

    query = request.args.get(
        "q", ""
    ).strip()

    if not query:
        return jsonify({
            "images": []
        })

    api_url = "https://commons.wikimedia.org/w/api.php"

    params = {
        "action": "query",
        "format": "json",
        "formatversion": "2",
        "generator": "search",
        "gsrsearch": query,
        "gsrnamespace": "6",
        "gsrlimit": "12",
        "prop": "imageinfo",
        "iiprop": "url",
        "iiurlwidth": "500",
        "origin": "*"
    }

    headers = {
        "User-Agent": "MedAI/1.0"
    }

    try:

        response = requests.get(
            api_url,
            params=params,
            headers=headers,
            timeout=20
        )

        response.raise_for_status()

        data = response.json()

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 502

    results = []

    for page in data.get("query", {}).get(
        "pages", []
    ):

        imageinfo = page.get(
            "imageinfo",
            []
        )

        if not imageinfo:
            continue

        info = imageinfo[0]

        url = (
            info.get("thumburl")
            or info.get("url")
        )

        if not url:
            continue

        results.append({
            "title": page.get(
                "title",
                "Image"
            ),
            "url": url,
            "page": (
                "https://commons.wikimedia.org/wiki/"
                + quote(
                    page.get("title", "")
                )
            )
        })

    return jsonify({
        "images": results
    })


# =========================================================
# IMAGE ANALYSIS / OCR
# =========================================================

@app.route("/analyze-image", methods=["POST"])
def analyze_image():

    image = request.files.get("image")

    mode = request.form.get(
        "mode",
        "analysis"
    )

    if not image:
        return safe_json_error(
            "Image is required."
        )

    try:

        data_url = file_to_data_url(image)

    except Exception as e:

        return safe_json_error(
            str(e)
        )

    if mode == "ocr":

        instruction = """
Read the image carefully and perform OCR.

Return the text you can read from the image.
Preserve paragraphs, headings, numbers and important formatting
as much as possible.

If text is unclear, say so.
Do not invent text.
"""

    else:

        instruction = """
Analyze this image carefully.

Describe important visible information.
If it appears medical, explain educationally what is visible,
but do not claim a definitive diagnosis from an image alone.

Mention uncertainty when appropriate.
"""

    messages = [
        {
            "role": "system",
            "content": BASE_SYSTEM
        },
        {
            "role": "user",
            "content": [
                {
                    "type": "text",
                    "text": instruction
                },
                {
                    "type": "image_url",
                    "image_url": {
                        "url": data_url
                    }
                }
            ]
        }
    ]

    answer, error = groq_request(
        messages,
        VISION_MODEL,
        0.2
    )

    if error:

        if isinstance(error, dict):
            return jsonify(error), error.get(
                "status_code", 500
            )

        return jsonify({
            "error": error
        }), 500

    return jsonify({
        "answer": answer
    })


# =========================================================
# CALCULATOR
# =========================================================

@app.route("/calculate", methods=["POST"])
def calculate():

    data = request.get_json(
        silent=True
    ) or {}

    expression = str(
        data.get("expression", "")
    ).strip()

    if not expression:
        return safe_json_error(
            "Expression is required."
        )

    try:

        result = calculate_expression(
            expression
        )

        return jsonify({
            "result": result
        })

    except Exception as e:

        return safe_json_error(
            "Invalid expression: " + str(e)
        )


# =========================================================
# WEB SEARCH
# =========================================================

@app.route("/search", methods=["GET"])
def web_search():

    query = request.args.get(
        "q", ""
    ).strip()

    if not query:
        return jsonify({
            "results": []
        })

    headers = {
        "User-Agent":
            "Mozilla/5.0 MedAI Search"
    }

    try:

        response = requests.get(
            "https://html.duckduckgo.com/html/",
            params={"q": query},
            headers=headers,
            timeout=20
        )

        response.raise_for_status()

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 502

    soup = BeautifulSoup(
        response.text,
        "html.parser"
    )

    results = []

    for item in soup.select(
        ".result"
    )[:10]:

        title_node = item.select_one(
            ".result__title"
        )

        link_node = item.select_one(
            ".result__a"
        )

        snippet_node = item.select_one(
            ".result__snippet"
        )

        if not link_node:
            continue

        href = link_node.get(
            "href"
        )

        title = link_node.get_text(
            " ",
            strip=True
        )

        snippet = (
            snippet_node.get_text(
                " ",
                strip=True
            )
            if snippet_node
            else ""
        )

        if href:

            results.append({
                "title": title,
                "url": href,
                "snippet": snippet
            })

    return jsonify({
        "results": results
    })


# =========================================================
# DOCUMENT EXTRACTION
# =========================================================

def extract_pdf(file_storage):

    reader = PdfReader(
        file_storage.stream
    )

    parts = []

    for page in reader.pages:

        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        if text:
            parts.append(text)

    return "\n\n".join(parts)


def extract_docx(file_storage):

    doc = Document(
        file_storage.stream
    )

    parts = []

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in doc.tables:

        for row in table.rows:

            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            parts.append(
                " | ".join(values)
            )

    return "\n".join(parts)


def extract_txt(file_storage):

    raw = file_storage.read()

    for encoding in [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1"
    ]:

        try:
            return raw.decode(
                encoding
            )
        except UnicodeDecodeError:
            continue

    return raw.decode(
        "utf-8",
        errors="replace"
    )


# =========================================================
# DOCUMENT ANALYSIS
# =========================================================

@app.route(
    "/analyze-document",
    methods=["POST"]
)
def analyze_document():

    file = request.files.get(
        "file"
    )

    instruction = request.form.get(
        "instruction",
        "Summarize this document."
    ).strip()

    if not file:
        return safe_json_error(
            "Document is required."
        )

    filename = (
        file.filename or ""
    ).lower()

    try:

        if filename.endswith(".pdf"):

            extracted = extract_pdf(
                file
            )

        elif filename.endswith(".docx"):

            extracted = extract_docx(
                file
            )

        elif filename.endswith(".txt"):

            extracted = extract_txt(
                file
            )

        else:

            return safe_json_error(
                "Supported files: PDF, DOCX, TXT."
            )

    except Exception as e:

        return jsonify({
            "error": "Could not read document.",
            "details": str(e)
        }), 400

    extracted = clean_text(
        extracted,
        45000
    )

    if not extracted:

        return safe_json_error(
            "No readable text found."
        )

    prompt = f"""
Instruction:
{instruction}

Document:
{extracted}
"""

    answer, error = groq_request([
        {
            "role": "system",
            "content": BASE_SYSTEM
        },
        {
            "role": "user",
            "content": prompt
        }
    ])

    if error:

        if isinstance(error, dict):
            return jsonify(error), error.get(
                "status_code", 500
            )

        return jsonify({
            "error": error
        }), 500

    return jsonify({
        "answer": answer
    })


# =========================================================
# DATA EXTRACTION
# =========================================================

def extract_csv(file_storage):

    raw = file_storage.read()

    text = raw.decode(
        "utf-8-sig",
        errors="replace"
    )

    reader = csv.reader(
        io.StringIO(text)
    )

    rows = list(reader)

    return rows


def extract_xlsx(file_storage):

    workbook = load_workbook(
        file_storage.stream,
        read_only=True,
        data_only=True
    )

    output = []

    for sheet in workbook.worksheets:

        output.append(
            f"\nSHEET: {sheet.title}"
        )

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = [
                "" if v is None
                else str(v)
                for v in row
            ]

            output.append(
                " | ".join(values)
            )

    return "\n".join(output)


# =========================================================
# DATA ANALYSIS
# =========================================================

@app.route(
    "/analyze-data",
    methods=["POST"]
)
def analyze_data():

    file = request.files.get(
        "file"
    )

    question = request.form.get(
        "question",
        "Analyze this dataset and provide useful insights."
    ).strip()

    if not file:
        return safe_json_error(
            "Data file is required."
        )

    filename = (
        file.filename or ""
    ).lower()

    try:

        if filename.endswith(".csv"):

            rows = extract_csv(
                file
            )

            lines = []

            for row in rows[:1000]:

                lines.append(
                    " | ".join(
                        str(x)
                        for x in row
                    )
                )

            extracted = "\n".join(
                lines
            )

        elif filename.endswith(
            ".xlsx"
        ):

            extracted = extract_xlsx(
                file
            )

        else:

            return safe_json_error(
                "Supported files: CSV and XLSX."
            )

    except Exception as e:

        return jsonify({
            "error": "Could not read data file.",
            "details": str(e)
        }), 400

    extracted = clean_text(
        extracted,
        40000
    )

    prompt = f"""
Question:
{question}

Dataset:
{extracted}

Give useful, understandable analysis.
Mention patterns, important numbers, possible issues,
and limitations where appropriate.
"""

    answer, error = groq_request([
        {
            "role": "system",
            "content": BASE_SYSTEM
        },
        {
            "role": "user",
            "content": prompt
        }
    ])

    if error:

        if isinstance(error, dict):
            return jsonify(error), error.get(
                "status_code", 500
            )

        return jsonify({
            "error": error
        }), 500

    return jsonify({
        "answer": answer
    })


# =========================================================
# URL ANALYSIS
# =========================================================

@app.route(
    "/analyze-url",
    methods=["POST"]
)
def analyze_url():

    data = request.get_json(
        silent=True
    ) or {}

    url = str(
        data.get("url", "")
    ).strip()

    question = str(
        data.get(
            "question",
            "Summarize this website."
        )
    ).strip()

    if not url:
        return safe_json_error(
            "URL is required."
        )

    parsed = urlparse(url)

    if parsed.scheme not in [
        "http",
        "https"
    ]:

        return safe_json_error(
            "Only HTTP and HTTPS URLs are allowed."
        )

    if not parsed.netloc:

        return safe_json_error(
            "Invalid URL."
        )

    hostname = (
        parsed.hostname or ""
    ).lower()

    blocked_hosts = [
        "localhost",
        "127.0.0.1",
        "0.0.0.0",
        "::1",
        "metadata.google.internal"
    ]

    if hostname in blocked_hosts:

        return safe_json_error(
            "This URL is not allowed."
        )

    try:

        response = requests.get(
            url,
            headers={
                "User-Agent":
                    "Mozilla/5.0 MedAI"
            },
            timeout=20,
            allow_redirects=True
        )

        response.raise_for_status()

    except Exception as e:

        return jsonify({
            "error": "Could not fetch URL.",
            "details": str(e)
        }), 502

    soup = BeautifulSoup(
        response.text,
        "html.parser"
    )

    for tag in soup([
        "script",
        "style",
        "noscript",
        "svg"
    ]):

        tag.decompose()

    title = (
        soup.title.get_text(
            " ",
            strip=True
        )
        if soup.title
        else ""
    )

    text = soup.get_text(
        " ",
        strip=True
    )

    text = clean_text(
        text,
        35000
    )

    prompt = f"""
Website title:
{title}

Website content:
{text}

User question:
{question}

Answer based only on the available website content.
If information is missing, say that it is not available.
"""

    answer, error = groq_request([
        {
            "role": "system",
            "content": BASE_SYSTEM
        },
        {
            "role": "user",
            "content": prompt
        }
    ])

    if error:

        if isinstance(error, dict):
            return jsonify(error), error.get(
                "status_code", 500
            )

        return jsonify({
            "error": error
        }), 500

    return jsonify({
        "answer": answer
    })


# =========================================================
# IMAGE GENERATION
# =========================================================

@app.route(
    "/generate-image",
    methods=["POST"]
)
def generate_image():

    data = request.get_json(
        silent=True
    ) or {}

    prompt = str(
        data.get("prompt", "")
    ).strip()

    if not prompt:
        return safe_json_error(
            "Prompt is required."
        )

    if len(prompt) > 1000:
        prompt = prompt[:1000]

    encoded_prompt = quote(
        prompt,
        safe=""
    )

    url = (
        "https://image.pollinations.ai/"
        "prompt/"
        + encoded_prompt
    )

    return jsonify({
        "url": url,
        "service": "Pollinations"
    })


# =========================================================
# CREATE DOCX
# =========================================================

@app.route(
    "/create-docx",
    methods=["POST"]
)
def create_docx():

    data = request.get_json(
        silent=True
    ) or {}

    title = str(
        data.get(
            "title",
            "MedAI Document"
        )
    ).strip()

    content = str(
        data.get(
            "content",
            ""
        )
    )

    if not content.strip():
        return safe_json_error(
            "Content is required."
        )

    document = Document()

    document.add_heading(
        title,
        level=1
    )

    paragraphs = content.split(
        "\n"
    )

    for paragraph in paragraphs:

        document.add_paragraph(
            paragraph
        )

    output = io.BytesIO()

    document.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="medai_document.docx",
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        )
    )


# =========================================================
# ERRORS
# =========================================================

@app.errorhandler(413)
def too_large(error):
    return jsonify({
        "error":
            "File too large. Maximum size is 12 MB."
    }), 413


@app.errorhandler(404)
def not_found(error):

    if request.path.startswith("/api"):
        return jsonify({
            "error": "Not found"
        }), 404

    return "Not found", 404


@app.errorhandler(500)
def server_error(error):

    return jsonify({
        "error": "Internal server error",
        "details": str(error)
    }), 500


# =========================================================
# LOCAL RUN
# =========================================================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=int(
            os.environ.get(
                "PORT",
                5000
            )
        ),
        debug=False
    )
